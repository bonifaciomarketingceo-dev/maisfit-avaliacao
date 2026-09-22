#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Robo do painel Mais Fit.
1) Baixa o dados.json da pasta PAINEL do Google Drive (fonte viva) -> ./dados.json
2) Remonta a planilha mestre (Excel) a partir do modelo + dados.json e grava
   de volta na pasta 2026 do Drive (mesmo arquivo/ID, so troca o conteudo).
A conta de servico (robo-painel) precisa ter acesso de Editor a pasta Avaliacoes.
Nao quebra o painel se a parte do Excel falhar (fica dentro de try/except).
"""
import os, io, json, sys, datetime

from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload

# ---- IDs das pastas/arquivos no Drive do Guilherme ----
PASTA_PAINEL = "102uAcp3wDATXU01a1bvKk14ctk9uJoNM"   # PAINEL - Dados e Programacao
PASTA_2026   = "1D_Nhhfk55g2x-8-pAfEaehEOMLw3TZXG"   # Avaliacoes/2026 (onde vive a mestre)
NOME_DADOS   = "dados.json"
NOME_MESTRE  = "Dashboard_Mais_Fit_Mestre_V2.xlsx"
MODELO_XLSX  = "modelo_planilha.xlsx"                 # template no repo
XLSX_MIME    = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SCOPES = ["https://www.googleapis.com/auth/drive"]


def drive_service():
    raw = os.environ.get("GOOGLE_SA_KEY", "").strip()
    if not raw:
        raise RuntimeError("Secret GOOGLE_SA_KEY nao encontrado.")
    info = json.loads(raw)
    creds = service_account.Credentials.from_service_account_info(info, scopes=SCOPES)
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def achar_arquivo(svc, pasta_id, nome):
    q = ("'%s' in parents and name='%s' and trashed=false" % (pasta_id, nome))
    r = svc.files().list(q=q, orderBy="modifiedTime desc",
                         fields="files(id,name,modifiedTime)",
                         pageSize=10, supportsAllDrives=True,
                         includeItemsFromAllDrives=True).execute()
    files = r.get("files", [])
    return files[0]["id"] if files else None


def baixar_dados(svc):
    fid = achar_arquivo(svc, PASTA_PAINEL, NOME_DADOS)
    if not fid:
        print("[dados] Nenhum dados.json na pasta PAINEL. Mantendo o do repo.")
        return False
    req = svc.files().get_media(fileId=fid)
    buf = io.BytesIO()
    dl = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    conteudo = buf.getvalue().decode("utf-8")
    json.loads(conteudo)  # valida
    with open("dados.json", "w", encoding="utf-8") as f:
        f.write(conteudo)
    print("[dados] dados.json baixado do Drive (%d bytes)." % len(conteudo))
    return True


# ------- Excel: remontar a aba Diario a partir do dados.json -------
COL = {"qtd": 4, "mat": 8, "novo": 9, "H": 10, "M": 11}
COL_PROF = {"Lucas": 12, "Edy": 13, "Eduardo": 14, "Luiz": 15, "Maria": 16}
COLS_LIMPAR = [4, 8, 9, 10, 11, 12, 13, 14, 15, 16]


def remontar_excel():
    import openpyxl
    dados = json.load(open("dados.json", encoding="utf-8"))
    meses = dados.get("meses", {})
    wb = openpyxl.load_workbook(MODELO_XLSX)
    ws = wb["Diário_Avaliações"]

    # mapa (mes,dia) -> linha
    linha_de = {}
    for r in range(2, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, datetime.datetime):
            linha_de[(a.month, a.day)] = r

    # limpa as colunas de dados de todas as linhas
    for r in linha_de.values():
        for c in COLS_LIMPAR:
            ws.cell(r, c).value = None

    total_lancado = 0
    for mstr, md in meses.items():
        m = int(mstr)
        dia = md.get("dia") or []
        det = md.get("det") or {}
        # quantidades por dia
        for i, q in enumerate(dia, start=1):
            if q:
                r = linha_de.get((m, i))
                if r:
                    ws.cell(r, COL["qtd"]).value = q
                    total_lancado += q
        # detalhe por dia
        for dstr, dd in det.items():
            r = linha_de.get((m, int(dstr)))
            if not r:
                continue
            if dd.get("q") is not None and not (dia and dstr and len(dia) >= int(dstr) and dia[int(dstr)-1]):
                ws.cell(r, COL["qtd"]).value = dd["q"]
            for k in ("mat", "novo", "H", "M"):
                if dd.get(k):
                    ws.cell(r, COL[k]).value = dd[k]
            for prof, n in (dd.get("profs") or {}).items():
                c = COL_PROF.get(prof)
                if c and n:
                    ws.cell(r, c).value = n
    wb.save(MODELO_XLSX + ".out.xlsx")
    print("[excel] Planilha remontada (total lancado ~%d)." % total_lancado)
    return MODELO_XLSX + ".out.xlsx"


def subir_excel(svc, caminho):
    media = MediaIoBaseUpload(io.FileIO(caminho, "rb"), mimetype=XLSX_MIME, resumable=True)
    fid = achar_arquivo(svc, PASTA_2026, NOME_MESTRE)
    if fid:
        svc.files().update(fileId=fid, media_body=media, supportsAllDrives=True).execute()
        print("[excel] Mestre atualizada no Drive (mesmo ID %s)." % fid)
    else:
        meta = {"name": NOME_MESTRE, "parents": [PASTA_2026]}
        novo = svc.files().create(body=meta, media_body=media,
                                  fields="id", supportsAllDrives=True).execute()
        print("[excel] Mestre criada no Drive (ID %s)." % novo["id"])


def main():
    svc = drive_service()
    baixar_dados(svc)          # painel (obrigatorio)
    try:                       # excel (best-effort)
        out = remontar_excel()
        subir_excel(svc, out)
    except Exception as e:
        print("[excel] AVISO: falhou remontar/subir o Excel: %r" % e)
    print("OK.")


if __name__ == "__main__":
    main()
